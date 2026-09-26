from pathlib import Path
from docx import Document
from pypdf import PdfReader
sources=[('mocks',r'C:\Users\Pavel\Downloads\Telegram Desktop\Описание_мок_экранов_Питомец_Финни_Грошик.docx'),('brief',r'C:\Users\Pavel\Downloads\Telegram Desktop\6_Департамент_финансов_города_Москвы.pdf'),('project',r'C:\Users\Pavel\Downloads\Лучше_Црезвоньте_Толу_2026.pdf')]
for name,path in sources:
    if path.endswith('.docx'):
        d=Document(path)
        text='\n'.join(p.text for p in d.paragraphs)
        for t in d.tables:
            text+='\n'+'\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows)
    else:
        d=PdfReader(path).pages
        text='\n'.join(f'\n--- PAGE {i+1} ---\n'+p.extract_text() for i,p in enumerate(d))
        print(name, 'pages',len(d))
    Path('work/'+name+'.txt').write_text(text,encoding='utf-8')
    print(name, 'chars',len(text))

